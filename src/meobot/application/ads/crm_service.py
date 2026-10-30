"""Doanh thu từ CRM (0083; ``docs/ads-report/CRM_REVENUE_PLAN.md``).

* :class:`AdsCrmSyncService` - ``ads.sync_crm_revenue`` (hourly, and on
  demand): the CRM's marketing users become :class:`AdsCrmAccount` rows (a new
  one takes its region from its name), then the payments of the days read
  **replace** what was stored for those days - a payment deleted or moved on
  the CRM goes too. An id seen on a payment but not listed is added
  (``in_directory`` false). The days read widen the synced range
  (``system_settings`` ``ads_crm``), inside which the CRM's numbers stand.
* :class:`AdsCrmService` - "Phòng Ads > Tài khoản CRM": each account's
  person / ◆ source / region, edited by hand or imported from the "Gắn mã"
  workbook (columns *CRM ID*, *Mã nhân viên APM*, *Hoặc tính vào nguồn ◆*,
  optional *Khu vực*, *Ghi chú*). A Trưởng phòng ADS, the OWNER / an ADMIN
  (``AdsAccess.edit_revenue``, else 403 ``ads_revenue_forbidden``).
"""

from __future__ import annotations

import base64
import binascii
import io
import unicodedata
import uuid
from collections import defaultdict
from collections.abc import Mapping
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from meobot.application.ads.access import AdsAccess
from meobot.application.ads.accounts_service import ads_people
from meobot.application.ads.crm_revenue import (
    SETTINGS_KEY,
    Coverage,
    account_subject,
    load_coverage,
)
from meobot.application.ads.revenue_service import AdsRevenueForbiddenError
from meobot.application.audit_service import AuditService
from meobot.core.errors import NotFoundError, ValidationError
from meobot.core.time import utcnow
from meobot.db.models.ads import AdsCrmAccount, AdsCrmPayment
from meobot.db.models.system_setting import SystemSetting
from meobot.db.models.user import User
from meobot.domain.ads.crm import (
    COUNTED_PAYMENT_TYPES,
    CrmRegion,
    HotlineKind,
    branch_site,
    hotline_kind_of,
    region_of,
)
from meobot.domain.ads.models import REVENUE_SOURCE_LABELS, RevenueSource
from meobot.domain.audit.models import AuditAction, AuditResult
from meobot.domain.identity.models import Actor
from meobot.integrations.crm.client import CrmClient, CrmError

#: The import's file limit.
MAX_IMPORT_BYTES = 5 * 1024 * 1024
#: "Không tính" in the source column = neither a person nor a source.
NOT_COUNTED = "khong tinh"
#: "Nhân viên đã nghỉ" in the source column = a former marketer (0084).
FORMER_WORDS = ("nhan vien da nghi", "da nghi", "nghi viec")


class AdsCrmInvalidError(ValidationError):
    code = "ads_crm_invalid"

    def __init__(self, message: str, *, field_name: str | None = None) -> None:
        super().__init__(message, details={"reason": "ads_crm_invalid", "field": field_name})


class AdsCrmImportError(ValidationError):
    code = "ads_crm_import_invalid"

    def __init__(self, message: str, problems: list[str]) -> None:
        super().__init__(
            message, details={"reason": "ads_crm_import_invalid", "problems": problems[:50]}
        )


class AdsCrmAccountNotFoundError(NotFoundError):
    code = "ads_crm_account_not_found"

    def __init__(self) -> None:
        super().__init__("Không tìm thấy tài khoản CRM.", details={"reason": self.code})


def crm_window(today: date) -> tuple[date, date]:
    """What the hourly run reads: this month so far; on the 1st-5th, from the
    1st of last month too (the CRM's books for it are still being closed)."""
    first = today.replace(day=1)
    if today.day <= 5:
        first = (first - timedelta(days=1)).replace(day=1)
    return first, today


# --- sync ---------------------------------------------------------------------------


@dataclass(slots=True)
class CrmSyncReport:
    since: date
    until: date
    users: int = 0
    payments: int = 0
    total: Decimal = Decimal(0)
    new_accounts: int = 0
    error: str | None = None
    configured: bool = True

    @property
    def status(self) -> str:
        return "FAILED" if self.error else "SUCCESS"

    @property
    def message(self) -> str:
        if not self.configured:
            return "Chưa cấu hình CRM_API_KEY: chưa lấy doanh thu từ CRM."
        head = f"Đồng bộ CRM {self.since:%d/%m} - {self.until:%d/%m}"
        if self.error:
            return f"{head}: {self.error}"
        total = f"{int(self.total):,}".replace(",", ".")
        return (
            f"{head}: {self.payments} khoản thu ({total} ₫), {self.users} tài khoản marketing"
            + (f", {self.new_accounts} tài khoản mới" if self.new_accounts else "")
            + "."
        )

    def as_task_result(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "message": self.message[:480],
            "payments": self.payments,
            "users": self.users,
        }


async def _save_coverage(session: AsyncSession, coverage: Coverage) -> None:
    _, row = await load_coverage(session)
    if row is None:
        session.add(
            SystemSetting(
                key=SETTINGS_KEY,
                value=coverage.as_value(),
                description="Phòng Ads: khoảng ngày đã lấy doanh thu từ CRM.",
            )
        )
    else:
        row.value = coverage.as_value()
        row.version = int(row.version) + 1


class AdsCrmSyncService:
    def __init__(self, session: AsyncSession, client: CrmClient) -> None:
        self._session = session
        self._client = client

    async def sync(self, since: date, until: date) -> CrmSyncReport:
        report = CrmSyncReport(since, until)
        coverage, _ = await load_coverage(self._session)
        now = utcnow()
        try:
            users = await self._client.marketing_users()
            branches = await self._client.branches()
            payments = await self._client.payments(since, until)
        except CrmError as error:
            report.error = error.message
            await _save_coverage(
                self._session,
                Coverage(
                    coverage.synced_from,
                    coverage.synced_to,
                    now,
                    report.status,
                    report.message,
                    coverage.branches,
                ),
            )
            return report
        accounts = {a.crm_id: a for a in await self._session.scalars(select(AdsCrmAccount))}
        for user in users:
            account = accounts.get(user.id)
            if account is None:
                account = AdsCrmAccount(crm_id=user.id, region=region_of(user.name))
                self._session.add(account)
                accounts[user.id] = account
                report.new_accounts += 1
            account.name = user.name[:200]
            account.active = user.active
            account.in_directory = True
            account.synced_at = now
        for payment in payments:
            if payment.marketing_id is None or payment.marketing_id in accounts:
                continue
            name = payment.marketing_name or f"CRM {payment.marketing_id}"
            account = AdsCrmAccount(
                crm_id=payment.marketing_id,
                name=name[:200],
                active=True,
                in_directory=False,
                region=region_of(name),
                synced_at=now,
            )
            self._session.add(account)
            accounts[payment.marketing_id] = account
            report.new_accounts += 1
        await self._session.flush()
        ids = [payment.id for payment in payments]
        stale = AdsCrmPayment.payment_date.between(since, until)
        await self._session.execute(
            delete(AdsCrmPayment).where(or_(stale, AdsCrmPayment.id.in_(ids)) if ids else stale)
        )
        self._session.add_all(
            AdsCrmPayment(
                id=payment.id,
                payment_date=payment.day,
                price=payment.price,
                payment_type=payment.payment_type,
                marketing_id=payment.marketing_id,
                branch_id=payment.branch_id,
                order_id=payment.order_id,
                is_upsale=payment.is_upsale,
                source_id=payment.source_id,
                synced_at=now,
            )
            for payment in {payment.id: payment for payment in payments}.values()
        )
        report.users = len(users)
        report.payments = len(payments)
        report.total = sum(
            (p.price for p in payments if p.payment_type in COUNTED_PAYMENT_TYPES), Decimal(0)
        )
        await _save_coverage(
            self._session,
            Coverage(
                min(filter(None, (coverage.synced_from, since))),
                max(filter(None, (coverage.synced_to, until))),
                now,
                report.status,
                report.message,
                {
                    branch.id: site
                    for branch in branches
                    if (site := branch_site(branch.name)) is not None
                }
                or coverage.branches,
            ),
        )
        await self._session.flush()
        return report


# --- the screen ---------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class PersonChoice:
    user_id: uuid.UUID
    name: str
    member_code: str | None
    in_ads: bool


@dataclass(frozen=True, slots=True)
class AccountView:
    account: AdsCrmAccount
    person: PersonChoice | None
    payments: int
    #: The month's counted thực thu.
    total: Decimal


@dataclass(frozen=True, slots=True)
class CrmOverview:
    month: date
    configured: bool
    coverage: Coverage
    accounts: list[AccountView]
    people: list[PersonChoice]
    #: The month's thực thu by who it counts for.
    mapped_total: Decimal
    unassigned_total: Decimal
    unassigned_accounts: int


@dataclass(slots=True)
class ImportResult:
    rows: int = 0
    changed: int = 0
    created: int = 0
    problems: list[str] = field(default_factory=list)


def _fold(text: object) -> str:
    raw = unicodedata.normalize("NFD", str(text or "").replace("đ", "d").replace("Đ", "D"))
    plain = "".join(c for c in raw if unicodedata.category(c) != "Mn").lower()
    return " ".join("".join(c if c.isalnum() else " " for c in plain).split())


_SOURCES_BY_LABEL = {_fold(label): source for source, label in REVENUE_SOURCE_LABELS.items()}
_REGIONS_BY_LABEL = {
    "hn": CrmRegion.HN,
    "ha noi": CrmRegion.HN,
    "hcm": CrmRegion.HCM,
    "sg": CrmRegion.HCM,
    "sai gon": CrmRegion.HCM,
    "ho chi minh": CrmRegion.HCM,
}


class AdsCrmService:
    def __init__(self, session: AsyncSession, audit: AuditService) -> None:
        self._session = session
        self._audit = audit

    @staticmethod
    def _require(access: AdsAccess) -> None:
        access.require_view()
        if not access.edit_revenue:
            raise AdsRevenueForbiddenError()

    async def _people(self) -> list[PersonChoice]:
        ads = {user.id for user in await ads_people(self._session)}
        users = await self._session.scalars(
            select(User).where(User.deleted_at.is_(None), User.active.is_(True))
        )
        choices = [
            PersonChoice(user.id, user.full_name, user.member_code, user.id in ads)
            for user in users
        ]
        return sorted(choices, key=lambda c: (not c.in_ads, c.name.casefold()))

    async def overview(self, access: AdsAccess, month: date, *, configured: bool) -> CrmOverview:
        self._require(access)
        first = month.replace(day=1)
        last = (first.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
        coverage, _ = await load_coverage(self._session)
        people = await self._people()
        by_id = {person.user_id: person for person in people}
        sums: dict[int, tuple[int, Decimal]] = defaultdict(lambda: (0, Decimal(0)))
        rows = await self._session.execute(
            select(
                AdsCrmPayment.marketing_id,
                func.count(AdsCrmPayment.id),
                func.sum(AdsCrmPayment.price),
            )
            .where(
                AdsCrmPayment.payment_date >= first,
                AdsCrmPayment.payment_date <= last,
                AdsCrmPayment.payment_type.in_(COUNTED_PAYMENT_TYPES),
                AdsCrmPayment.marketing_id.is_not(None),
            )
            .group_by(AdsCrmPayment.marketing_id)
        )
        for marketing_id, count, total in rows:
            sums[marketing_id] = (int(count), Decimal(total or 0))
        views = []
        for account in await self._session.scalars(select(AdsCrmAccount)):
            count, total = sums.get(account.crm_id, (0, Decimal(0)))
            person = by_id.get(account.user_id) if account.user_id else None
            views.append(AccountView(account, person, count, total))
        views.sort(
            key=lambda v: (
                not v.account.active,
                -v.total,
                v.account.name.casefold(),
            )
        )
        mapped = [v for v in views if account_subject(v.account)]
        loose = [v for v in views if not account_subject(v.account) and v.total]
        return CrmOverview(
            month=first,
            configured=configured,
            coverage=coverage,
            accounts=views,
            people=people,
            mapped_total=sum((v.total for v in mapped), Decimal(0)),
            unassigned_total=sum((v.total for v in loose), Decimal(0)),
            unassigned_accounts=len(loose),
        )

    async def update(
        self,
        *,
        actor: Actor,
        access: AdsAccess,
        request_id: uuid.UUID,
        crm_id: int,
        fields: Mapping[str, Any],
    ) -> AdsCrmAccount:
        """``user_id`` / ``source`` / ``former`` (one or none) / ``region`` /
        ``hotline_kind`` / ``note``; a key left out keeps its value."""
        self._require(access)
        account = await self._session.get(AdsCrmAccount, crm_id)
        if account is None:
            raise AdsCrmAccountNotFoundError()
        before = self._snapshot(account)
        if "user_id" in fields and fields["user_id"] is not None:
            user = await self._session.get(User, fields["user_id"])
            if user is None or user.deleted_at is not None:
                raise AdsCrmInvalidError("Người được chọn không tồn tại.", field_name="user_id")
        user_id = fields.get("user_id", account.user_id)
        source = fields.get("source", account.source)
        if user_id is not None and source is not None:
            if "user_id" in fields and "source" in fields:
                raise AdsCrmInvalidError(
                    "Một tài khoản CRM thuộc một người hoặc một nguồn, không cả hai.",
                    field_name="source",
                )
            # Picking one drops the other.
            if "user_id" in fields:
                source = None
            else:
                user_id = None
        former = bool(fields.get("former", account.former))
        if former and (user_id is not None or source is not None):
            if fields.get("former"):
                if ("user_id" in fields and fields["user_id"] is not None) or (
                    "source" in fields and fields["source"] is not None
                ):
                    raise AdsCrmInvalidError(
                        "Nhân viên đã nghỉ không gắn thêm người hay nguồn.",
                        field_name="former",
                    )
                user_id = source = None
            else:
                former = False
        account.user_id = user_id
        account.source = source
        account.former = former
        if source is not RevenueSource.HOTLINE:
            account.hotline_kind = None
        elif "hotline_kind" in fields:
            account.hotline_kind = fields["hotline_kind"]
        elif account.hotline_kind is None:
            account.hotline_kind = hotline_kind_of(account.name)
        if "region" in fields:
            account.region = fields["region"]
        if "note" in fields:
            note = (fields["note"] or "").strip()
            account.note = note[:500] or None
        account.updated_by_user_id = actor.user_id
        account.updated_at = utcnow()
        await self._session.flush()
        await self._audit.record_action(
            request_id=request_id,
            actor=actor,
            action=AuditAction.ADS_CRM_ACCOUNT_UPDATED.value,
            result=AuditResult.SUCCESS,
            entity_type="ads_crm_account",
            entity_id=str(crm_id),
            before_data=before,
            after_data=self._snapshot(account),
        )
        return account

    @staticmethod
    def _snapshot(account: AdsCrmAccount) -> dict[str, object]:
        return {
            "user_id": str(account.user_id) if account.user_id else None,
            "source": account.source.value if account.source else None,
            "region": account.region.value if account.region else None,
            "hotline_kind": account.hotline_kind.value if account.hotline_kind else None,
            "former": account.former,
            "note": account.note,
        }

    async def import_mapping(
        self, *, actor: Actor, access: AdsAccess, request_id: uuid.UUID, data: str
    ) -> ImportResult:
        """The "Gắn mã" workbook: every row with a CRM ID sets that account's
        person (by *Mã nhân viên APM*) or ◆ source - blank both = not counted.
        Any problem refuses the whole file (422 ``ads_crm_import_invalid``)."""
        self._require(access)
        try:
            raw = base64.b64decode(data, validate=True)
        except (binascii.Error, ValueError) as error:
            raise AdsCrmInvalidError("File không đọc được.", field_name="data") from error
        if len(raw) > MAX_IMPORT_BYTES:
            raise AdsCrmInvalidError("File lớn hơn 5 MB.", field_name="data")
        try:
            book = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        except Exception as error:  # openpyxl raises several kinds
            raise AdsCrmInvalidError(
                "File không phải Excel (.xlsx) hợp lệ.", field_name="data"
            ) from error
        codes = {
            (user.member_code or "").upper(): user
            for user in await self._session.scalars(
                select(User).where(User.deleted_at.is_(None), User.member_code.is_not(None))
            )
        }
        accounts = {a.crm_id: a for a in await self._session.scalars(select(AdsCrmAccount))}
        result = ImportResult()
        planned: list[
            tuple[
                int,
                str | None,
                uuid.UUID | None,
                RevenueSource | None,
                Any,
                bool,
                HotlineKind | None,
            ]
        ] = []
        found = False
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            columns: dict[str, int] = {}
            for number, row in enumerate(rows, 1):
                cells = list(row)
                if not columns:
                    folded = [_fold(cell) for cell in cells]
                    if "crm id" in folded and any(f.startswith("ma nhan vien") for f in folded):
                        for index, name in enumerate(folded):
                            if name == "crm id":
                                columns["id"] = index
                            elif name.startswith("ma nhan vien"):
                                columns["code"] = index
                            elif name.startswith("hoac tinh vao nguon") or name == "nguon":
                                columns["source"] = index
                            elif name.startswith("khu vuc"):
                                columns["region"] = index
                            elif name.startswith("ten marketing") or name == "ten tren crm":
                                columns["name"] = index
                            elif name == "ghi chu":
                                columns["note"] = index
                    continue

                def cell(
                    key: str, values: list[Any] = cells, where: dict[str, int] = columns
                ) -> Any:
                    index = where.get(key)
                    return values[index] if index is not None and index < len(values) else None

                raw_id = cell("id")
                if raw_id in (None, ""):
                    continue
                found = True
                where = f"{sheet.title} dòng {number}"
                try:
                    crm_id = int(str(raw_id).strip().split(".")[0])
                except ValueError:
                    result.problems.append(f"{where}: CRM ID “{raw_id}” không phải số.")
                    continue
                result.rows += 1
                code = str(cell("code") or "").strip().upper()
                source_text = _fold(cell("source"))
                user_id: uuid.UUID | None = None
                source: RevenueSource | None = None
                if code:
                    user = codes.get(code)
                    if user is None:
                        result.problems.append(f"{where}: mã nhân viên “{code}” không có trên APM.")
                        continue
                    user_id = user.id
                former = source_text in FORMER_WORDS
                kind: HotlineKind | None = None
                if former and user_id is not None:
                    result.problems.append(f"{where}: nhân viên đã nghỉ thì để trống mã nhân viên.")
                    continue
                if source_text and source_text != NOT_COUNTED and not former:
                    if source_text.startswith("hotline"):
                        source, kind = RevenueSource.HOTLINE, hotline_kind_of(source_text)
                    else:
                        source = _SOURCES_BY_LABEL.get(source_text)
                    if source is None:
                        result.problems.append(
                            f"{where}: nguồn “{cell('source')}” không có trong danh sách."
                        )
                        continue
                    if user_id is not None:
                        result.problems.append(
                            f"{where}: chỉ điền mã nhân viên HOẶC nguồn ◆, không cả hai."
                        )
                        continue
                region_text = _fold(cell("region"))
                region: Any = ...
                if region_text:
                    region = _REGIONS_BY_LABEL.get(region_text)
                    if region is None:
                        result.problems.append(
                            f"{where}: khu vực “{cell('region')}” phải là HN hoặc HCM."
                        )
                        continue
                crm_name = str(cell("name") or "").strip() or None
                planned.append((crm_id, crm_name, user_id, source, region, former, kind))
        if not found:
            raise AdsCrmImportError(
                "Không thấy bảng có cột “CRM ID” và “Mã nhân viên APM”.",
                ["Dùng file “Gắn mã” (sheet có hàng tiêu đề CRM ID, Mã nhân viên APM)."],
            )
        if result.problems:
            raise AdsCrmImportError(
                f"File có {len(result.problems)} dòng lỗi, chưa lưu gì.", result.problems
            )
        now = utcnow()
        for crm_id, crm_name, user_id, source, region, former, kind in planned:
            account = accounts.get(crm_id)
            if account is None:
                label = crm_name or f"CRM {crm_id}"
                account = AdsCrmAccount(
                    crm_id=crm_id,
                    name=label[:200],
                    active=True,
                    in_directory=False,
                    region=region_of(label),
                    synced_at=now,
                )
                self._session.add(account)
                accounts[crm_id] = account
                result.created += 1
            before = (
                account.user_id,
                account.source,
                account.region,
                account.former,
                account.hotline_kind,
            )
            account.user_id = user_id
            account.source = source
            account.former = former
            account.hotline_kind = (
                (kind or account.hotline_kind or hotline_kind_of(account.name))
                if source is RevenueSource.HOTLINE
                else None
            )
            if region is not ...:
                account.region = region
            if before != (
                account.user_id,
                account.source,
                account.region,
                account.former,
                account.hotline_kind,
            ):
                account.updated_by_user_id = actor.user_id
                account.updated_at = now
                result.changed += 1
        await self._session.flush()
        await self._audit.record_action(
            request_id=request_id,
            actor=actor,
            action=AuditAction.ADS_CRM_IMPORTED.value,
            result=AuditResult.SUCCESS,
            entity_type="ads_crm_account",
            entity_id="import",
            after_data={"rows": result.rows, "changed": result.changed, "created": result.created},
        )
        return result


__all__ = [
    "AdsCrmAccountNotFoundError",
    "AdsCrmImportError",
    "AdsCrmInvalidError",
    "AdsCrmService",
    "AdsCrmSyncService",
    "CrmOverview",
    "CrmSyncReport",
    "ImportResult",
    "crm_window",
]
