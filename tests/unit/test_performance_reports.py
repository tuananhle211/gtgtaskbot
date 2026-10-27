"""The points reports ("Hiệu suất", ``docs/perf/PERFORMANCE_REPORTS_PLAN.md``).

* everybody sees their own month: points, steps, the lines, what is still
  under way (dự kiến, not in the total);
* a Trưởng phòng sees only their direct reports ("Trưởng quản lý" = them);
* the company ("Toàn công ty", the Excel file): the OWNER, an ADMIN, or a
  ``PERFORMANCE`` grant - nobody else by default;
* the points count for the department of the step's phase.
"""

from __future__ import annotations

import io
import uuid
from typing import Any

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update

from meobot.db.models.department import Department, DepartmentMember
from meobot.db.models.order import OrderNodePoints
from meobot.db.models.user import User
from meobot.domain.identity.models import Role
from meobot.domain.org.models import DepartmentFunction
from tests.unit.pr_world import World
from tests.unit.test_design_ban import do
from tests.unit.test_order_commands import Ads, person
from tests.unit.test_task_codes import at_dung, coded_world
from tests.unit.test_unified_permissions import give

pytestmark = pytest.mark.asyncio


async def dung_done(world: World, ads: Ads) -> str:
    """A Biên kịch > Dựng order done by the writer (N3 120) and the editor
    (D3 240); returns its id."""
    code, detail = await at_dung(world, ads)
    done = do(
        world,
        ads.editor,
        code,
        "ads:SUBMIT_WORK:",
        link="https://e/v",
        task_codes=[{"code_id": world.task_codes["D3"]}],  # type: ignore[attr-defined]
    )
    assert done.status_code == 200, done.json()
    return str(detail["order"]["id"])


async def report_to(world: World, user_id: uuid.UUID, head_id: uuid.UUID) -> None:
    await world.session.execute(
        update(DepartmentMember)
        .where(DepartmentMember.user_id == user_id, DepartmentMember.left_at.is_(None))
        .values(manager_user_id=head_id)
    )
    await world.session.flush()


def get(world: World, user: User, path: str) -> Any:
    world.act_as(user)
    return world.client.get(path)


async def test_01_everybody_sees_their_own_month(world: World) -> None:
    ads, _ = await coded_world(world)
    await dung_done(world, ads)
    # A second order waits at Dựng: the editor has it, not yet handed in.
    await at_dung(world, ads)

    world.act_as(ads.editor)
    body = world.client.get("/api/performance/people/me").json()
    assert body["name"] == ads.editor.full_name
    figures = body["figures"]
    assert figures["points"] == 240 and figures["steps_done"] == 1
    # Dự kiến: the step under way, beside the total, never in it.
    assert figures["pending_points"] == 240 and figures["pending_count"] == 1
    assert body["by_phase"] == {"DUNG": 240} and sum(body["by_week"]) == 240
    statuses = sorted(line["status"] for line in body["lines"])
    assert statuses == ["EARNED", "IN_PROGRESS"]
    earned = next(line for line in body["lines"] if line["status"] == "EARNED")
    assert earned["phase_label"] == "Dựng" and earned["items"][0]["code"] == "D3"
    assert earned["items"][0]["points"] == 240
    # Their own count of each task code.
    assert [(row["code"], row["quantity"], row["points"]) for row in body["codes"]] == [
        ("D3", 1, 240)
    ]

    # Nobody else's, and no team / company screen for a plain member.
    assert get(world, ads.editor, f"/api/performance/people/{ads.writer.id}").status_code == 403
    for path in ("/api/performance/team", "/api/performance/overview"):
        refused = get(world, ads.editor, path)
        assert refused.status_code == 403
        assert refused.json()["code"] == "performance_forbidden"
    flags = get(world, ads.editor, "/api/units/me").json()["performance"]
    assert flags == {
        "view_team": False,
        "view_all": False,
        "export": False,
        "manage_coefficients": False,
    }


async def test_02_points_count_for_the_department_of_the_phase(world: World) -> None:
    ads, _ = await coded_world(world)
    order_id = await dung_done(world, ads)
    rows = (
        await world.session.execute(
            select(OrderNodePoints.user_id, Department.function)
            .join(Department, Department.id == OrderNodePoints.department_id)
            .where(OrderNodePoints.order_id == uuid.UUID(order_id))
        )
    ).all()
    assert dict(rows) == {
        ads.writer.id: DepartmentFunction.BIEN_TAP,
        ads.editor.id: DepartmentFunction.DUNG,
    }


async def test_03_a_trưởng_phòng_sees_only_their_direct_reports(world: World) -> None:
    ads, _ = await coded_world(world)
    await dung_done(world, ads)
    # Heading the Dựng department is not enough: the editor reports to nobody.
    assert get(world, ads.lead_dung, f"/api/performance/people/{ads.editor.id}").status_code == 403
    team = get(world, ads.lead_dung, "/api/performance/team").json()
    assert team["rows"] == [] and [item["name"] for item in team["departments"]]
    assert get(world, ads.lead_dung, "/api/units/me").json()["performance"]["view_team"] is True

    await report_to(world, ads.editor.id, ads.lead_dung.id)
    mine = get(world, ads.lead_dung, f"/api/performance/people/{ads.editor.id}")
    assert mine.status_code == 200 and mine.json()["figures"]["points"] == 240
    team = get(world, ads.lead_dung, "/api/performance/team").json()
    assert [row["name"] for row in team["rows"]] == [ads.editor.full_name]
    assert team["totals"]["points"] == 240 and team["view_all"] is False
    # Never someone else's people, never the company.
    assert get(world, ads.lead_dung, f"/api/performance/people/{ads.writer.id}").status_code == 403
    assert get(world, ads.lead_dung, "/api/performance/overview").status_code == 403
    assert get(world, ads.lead_dung, "/api/performance/export.xlsx").status_code == 403
    # Only a department they head narrows the table.
    other = (
        await world.session.scalars(
            select(Department.id).where(Department.function == DepartmentFunction.BIEN_TAP)
        )
    ).first()
    narrowed = get(world, ads.lead_dung, f"/api/performance/team?department_id={other}")
    assert narrowed.status_code == 403


async def test_04_the_owner_and_an_admin_see_the_company_and_export(world: World) -> None:
    ads, _ = await coded_world(world)
    await dung_done(world, ads)
    admin = await person(world, "Quản trị", Role.ADMIN)
    for viewer in (world.owner, admin):
        overview = get(world, viewer, "/api/performance/overview")
        assert overview.status_code == 200, overview.json()
        body = overview.json()
        assert body["points"] == 360 and body["steps_done"] == 2 and body["people"] == 2
        assert sorted(row["points"] for row in body["departments"]) == [120, 240]
        assert body["top"][0]["name"] == ads.editor.full_name
        assert {row["code"] for row in body["codes"]} == {"N3", "D3"}
        assert body["trend"][-1]["points"] == 360 and len(body["trend"]) == 6
        assert body["export"] is True
    # Any person's month, a department's table.
    assert get(world, admin, f"/api/performance/people/{ads.writer.id}").status_code == 200
    dung = (
        await world.session.scalars(
            select(Department.id).where(Department.function == DepartmentFunction.DUNG)
        )
    ).first()
    team = get(world, admin, f"/api/performance/team?department_id={dung}").json()
    assert ads.editor.full_name in [row["name"] for row in team["rows"]]
    assert ads.writer.full_name not in [row["name"] for row in team["rows"]]
    assert get(world, admin, f"/api/performance/people/{uuid.uuid4()}").status_code == 404

    sheet = get(world, world.owner, "/api/performance/export.xlsx")
    assert sheet.status_code == 200
    assert "diem-hieu-suat-" in sheet.headers["content-disposition"]
    book = load_workbook(io.BytesIO(sheet.content))
    assert book.sheetnames == ["Tổng hợp", "Chi tiết", "Theo mã task"]
    names = [row[0].value for row in book["Tổng hợp"].iter_rows(min_row=3)]
    assert ads.editor.full_name in names and names[-1] == "Tổng"
    details = list(book["Chi tiết"].iter_rows(min_row=2, values_only=True))
    assert {row[6] for row in details} == {"N3", "D3"}


async def test_05_a_performance_grant_opens_the_company(world: World) -> None:
    ads, _ = await coded_world(world)
    await dung_done(world, ads)
    assert get(world, ads.orderer, "/api/performance/overview").status_code == 403
    made = give(
        world, world.owner, ads.orderer, module="PERFORMANCE", permission="VIEW", scope="ALL"
    )
    assert made.status_code == 201, made.json()
    assert get(world, ads.orderer, "/api/performance/overview").status_code == 200
    assert get(world, ads.orderer, f"/api/performance/people/{ads.editor.id}").status_code == 200
    flags = get(world, ads.orderer, "/api/units/me").json()
    assert flags["performance"] == {
        "view_team": True,
        "view_all": True,
        "export": False,
        "manage_coefficients": False,
    }
    assert flags["can_view_profiles"] is True
    profile = get(world, ads.orderer, f"/api/account/members/{ads.editor.id}/profile")
    assert profile.status_code == 200
    # Exporting is its own grant.
    assert get(world, ads.orderer, "/api/performance/export.xlsx").status_code == 403
    give(world, world.owner, ads.orderer, module="PERFORMANCE", permission="EXPORT", scope="ALL")
    assert get(world, ads.orderer, "/api/performance/export.xlsx").status_code == 200


# --- "Hệ số điểm" (0076) ---------------------------------------------------------------


def next_month(label: str) -> str:
    year, month = (int(part) for part in label.split("-"))
    year, month = (year + 1, 1) if month == 12 else (year, month + 1)
    return f"{year:04d}-{month:02d}"


def put_coefficient(world: World, user: User, target: User, **body: Any) -> Any:
    world.act_as(user)
    return world.client.put(f"/api/performance/coefficients/{target.id}", json=body)


async def test_06_a_coefficient_multiplies_the_month_s_points_everywhere(world: World) -> None:
    ads, _ = await coded_world(world)
    await dung_done(world, ads)
    month = get(world, ads.editor, "/api/performance/people/me").json()["month"]
    # Default 1: the catalogue's points.
    assert get(world, ads.editor, "/api/performance/people/me").json()["figures"]["points"] == 240

    made = put_coefficient(
        world, world.owner, ads.editor, coefficient="1,5", month=month, note="Nhân sự chính"
    )
    assert made.status_code == 200, made.json()
    assert made.json()["coefficient"] == "1.50" and made.json()["effective_month"] == month

    body = get(world, ads.editor, "/api/performance/people/me").json()
    # Only the points after the coefficient are shown: lines, items, totals.
    assert body["figures"]["points"] == 360
    earned = next(line for line in body["lines"] if line["status"] == "EARNED")
    assert earned["points"] == 360 and earned["items"][0]["points"] == 360
    assert body["by_phase"] == {"DUNG": 360} and sum(body["by_week"]) == 360
    overview = get(world, world.owner, "/api/performance/overview").json()
    assert overview["points"] == 480  # writer 120 x 1 + editor 240 x 1.5
    assert {row["code"]: row["points"] for row in overview["codes"]} == {"N3": 120, "D3": 360}
    assert overview["trend"][-1]["points"] == 480
    profile = get(world, world.owner, f"/api/account/members/{ads.editor.id}/profile").json()
    assert float(profile["stats"]["points"]) == 360
    sheet = get(world, world.owner, "/api/performance/export.xlsx")
    book = load_workbook(io.BytesIO(sheet.content))
    details = {row[6]: row[10] for row in book["Chi tiết"].iter_rows(min_row=2, values_only=True)}
    assert details == {"N3": 120, "D3": 360}


async def test_07_a_coefficient_applies_from_its_month_on(world: World) -> None:
    ads, _ = await coded_world(world)
    await dung_done(world, ads)
    month = get(world, ads.editor, "/api/performance/people/me").json()["month"]
    later = next_month(month)
    assert (
        put_coefficient(world, world.owner, ads.editor, coefficient=2, month=later).status_code
        == 200
    )
    # This month keeps 1; the table shows the change to come.
    assert get(world, ads.editor, "/api/performance/people/me").json()["figures"]["points"] == 240
    world.act_as(world.owner)
    rows = world.client.get(f"/api/performance/coefficients?month={month}").json()["rows"]
    editor = next(row for row in rows if row["user_id"] == str(ads.editor.id))
    assert editor["coefficient"] == "1" and editor["effective_month"] is None
    assert editor["next_coefficient"] == "2.00" and editor["next_month"] == later
    rows = world.client.get(f"/api/performance/coefficients?month={later}").json()["rows"]
    editor = next(row for row in rows if row["user_id"] == str(ads.editor.id))
    assert editor["coefficient"] == "2.00" and editor["effective_month"] == later
    # A second change for this month supersedes nothing earlier: it is history.
    put_coefficient(world, world.owner, ads.editor, coefficient="0.8", month=month)
    assert get(world, ads.editor, "/api/performance/people/me").json()["figures"]["points"] == 192
    world.act_as(world.owner)
    history = world.client.get(f"/api/performance/coefficients/{ads.editor.id}").json()
    assert [(item["effective_month"], item["coefficient"]) for item in history] == [
        (later, "2.00"),
        (month, "0.80"),
    ]
    assert history[0]["created_by"] == world.owner.full_name


async def test_08_only_the_owner_sets_coefficients(world: World) -> None:
    ads, _ = await coded_world(world)
    admin = await person(world, "Quản trị", Role.ADMIN)
    month = get(world, ads.editor, "/api/performance/people/me").json()["month"]
    assert get(world, world.owner, "/api/units/me").json()["performance"]["manage_coefficients"]
    for user in (admin, ads.lead_dung, ads.editor):
        assert (
            get(world, user, "/api/units/me").json()["performance"]["manage_coefficients"] is False
        )
        refused = put_coefficient(world, user, ads.writer, coefficient=2, month=month)
        assert refused.status_code == 403
        assert refused.json()["code"] == "performance_coefficient_forbidden"
        assert get(world, user, "/api/performance/coefficients").status_code == 403
    # Nobody but the OWNER may grant it - not even an ADMIN.
    by_admin = give(
        world, admin, ads.orderer, module="PERFORMANCE", permission="COEFFICIENT", scope="ALL"
    )
    assert by_admin.status_code == 422
    assert by_admin.json()["error"]["details"]["reason"] == "permission_not_grantable"
    by_owner = give(
        world, world.owner, ads.orderer, module="PERFORMANCE", permission="COEFFICIENT", scope="ALL"
    )
    assert by_owner.status_code == 201, by_owner.json()
    assert (
        put_coefficient(world, ads.orderer, ads.writer, coefficient=2, month=month).status_code
        == 200
    )


@pytest.mark.parametrize("value", ["0", "10.5", "1.234", "abc", "-1"])
async def test_09_a_coefficient_is_0_01_to_10_with_two_decimals(world: World, value: str) -> None:
    ads, _ = await coded_world(world)
    refused = put_coefficient(world, world.owner, ads.editor, coefficient=value, month="2026-10")
    assert refused.status_code == 422, refused.json()
    assert refused.json()["details"]["reason"] == "coefficient_invalid"


async def test_10_a_department_s_report_counts_each_task_code(world: World) -> None:
    ads, _ = await coded_world(world)
    await dung_done(world, ads)
    await dung_done(world, ads)
    await report_to(world, ads.editor.id, ads.lead_dung.id)
    # The head: their reports' codes - two D3 by one editor.
    team = get(world, ads.lead_dung, "/api/performance/team").json()
    assert [
        (row["phase_label"], row["code"], row["form_label"], row["quantity"], row["people"])
        for row in team["codes"]
    ] == [("Dựng", "D3", None, 2, 1)]
    assert team["codes"][0]["points"] == 480 and team["codes"][0]["steps"] == 2

    # The company viewer, by department: only the points counted there.
    dung = (
        await world.session.scalars(
            select(Department.id).where(Department.function == DepartmentFunction.DUNG)
        )
    ).first()
    everybody = get(world, world.owner, "/api/performance/team").json()["codes"]
    assert [(row["code"], row["quantity"]) for row in everybody] == [("N3", 2), ("D3", 2)]
    narrowed = get(world, world.owner, f"/api/performance/team?department_id={dung}").json()
    assert [(row["code"], row["quantity"]) for row in narrowed["codes"]] == [("D3", 2)]

    # After the coefficient, like every other figure; in the Excel file too.
    month = team["month"]
    put_coefficient(world, world.owner, ads.editor, coefficient="1.5", month=month)
    team = get(world, ads.lead_dung, "/api/performance/team").json()
    assert team["codes"][0]["points"] == 720
    sheet = get(world, world.owner, f"/api/performance/export.xlsx?department_id={dung}")
    rows = list(
        load_workbook(io.BytesIO(sheet.content))["Theo mã task"].iter_rows(values_only=True)
    )
    assert rows[0][:5] == ("Bước", "Mã task", "Tên / ví dụ", "Hình thức", "Số lượng")
    assert rows[1][0:2] == ("Dựng", "D3") and rows[1][3] == "Làm mới" and rows[1][4:] == (2, 1, 720)
    assert rows[-1][0] == "Tổng" and rows[-1][6] == 720


async def test_11_the_gantt_shows_each_person_s_steps_in_the_month(world: World) -> None:
    ads, _ = await coded_world(world)
    await dung_done(world, ads)
    await at_dung(world, ads)  # a second one, taken, not handed in

    mine = get(world, ads.editor, "/api/performance/timeline").json()
    assert [group["name"] for group in mine["groups"]] == [ads.editor.full_name]
    bars = mine["groups"][0]["bars"]
    done = next(bar for bar in bars if bar["status"] == "EARNED")
    open_ = next(bar for bar in bars if bar["status"] == "IN_PROGRESS")
    assert done["phase"] == "DUNG" and done["codes"] == ["D3"] and done["points"] == 240
    assert done["done_at"] is not None and done["start"] <= done["done_at"]
    assert open_["done_at"] is None and open_["points"] == 240 and open_["late"] is False
    assert mine["groups"][0]["points"] == 240 and mine["groups"][0]["pending_points"] == 240
    # A plain member sees only their own; no department view.
    refused = get(world, ads.editor, f"/api/performance/timeline?department_id={uuid.uuid4()}")
    assert refused.status_code == 403

    # Every Trưởng phòng sees every department (2026-10-27): everybody with a
    # step in the month, or one department's people.
    head = get(world, ads.lead_dung, "/api/performance/timeline").json()
    assert {group["name"] for group in head["groups"]} == {
        ads.writer.full_name,
        ads.editor.full_name,
    }
    assert head["view_all"] is True and head["departments"]
    bien_tap = (
        await world.session.scalars(
            select(Department.id).where(Department.function == DepartmentFunction.BIEN_TAP)
        )
    ).first()
    one = get(world, ads.lead_dung, f"/api/performance/timeline?department_id={bien_tap}")
    assert one.status_code == 200
    assert [g["name"] for g in one.json()["groups"] if g["bars"]] == [ads.writer.full_name]
    # The points reports still keep them to their own people.
    assert get(world, ads.lead_dung, f"/api/performance/people/{ads.writer.id}").status_code == 403
    # The company: only people with a step in the month, points after coefficients.
    put_coefficient(world, world.owner, ads.editor, coefficient="2", month=mine["month"])
    company = get(world, world.owner, "/api/performance/timeline").json()
    groups = {group["name"]: group for group in company["groups"]}
    assert set(groups) == {ads.writer.full_name, ads.editor.full_name}
    assert groups[ads.editor.full_name]["points"] == 480
    assert groups[ads.writer.full_name]["bars"][0]["phase_label"] == "Biên kịch"


async def test_12_the_gantt_shows_what_is_ahead(world: World) -> None:
    from tests.unit.test_order_commands import create

    ads, _ = await coded_world(world)
    await at_dung(world, ads)  # open, taken now
    # Handed to the editor ahead; the order is still at Biên kịch.
    ahead = create(
        world,
        ads,
        video_type="BD",
        preassigned={"BIEN_TAP": str(ads.writer.id), "DUNG": str(ads.editor.id)},
    )
    mine = get(world, ads.editor, "/api/performance/timeline").json()
    group = mine["groups"][0]
    upcoming = [bar for bar in group["bars"] if bar["status"] == "UPCOMING"]
    assert [bar["order_code"] for bar in upcoming] == [ahead["order"]["code"]]
    # Its deadline is the order's wish while the step has none; not in "dự kiến".
    assert upcoming[0]["deadline_at"].startswith("2030-01-01")
    assert upcoming[0]["codes"] == ["D3"] and upcoming[0]["points"] == 240
    assert group["upcoming"] == 1 and group["pending_points"] == 240

    # A later month still shows what runs into it (open steps up to their
    # deadline, steps ahead up to theirs); a past month shows nothing ahead.
    later = next_month(next_month(mine["month"]))
    ahead_view = get(world, ads.editor, f"/api/performance/timeline?month={later}").json()
    statuses = sorted(bar["status"] for bar in ahead_view["groups"][0]["bars"])
    assert "UPCOMING" in statuses
    past = get(world, ads.editor, "/api/performance/timeline?month=2026-01").json()
    assert past["groups"][0]["bars"] == []
